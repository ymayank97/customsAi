"""Excel I/O for the BE upload + annotated output."""
from __future__ import annotations

import pathlib
import math
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Column synonyms for input mapping (case-insensitive substring match)
DESC_KEYS = ["item desc", "item description", "description", "goods description"]
CTH_KEYS = ["cth", "rita cth", "hs code", "hsn"]
BCD_KEYS = ["bcd rate", "bcd %", "bcd"]
NTFN_KEYS = ["bcd notn", "notification", "notn"]
ASSESS_KEYS = ["assess val", "assessable value", "assess value", "value"]
IMPORTER_KEYS = ["importer", "iec name", "consignee"]
BE_NO_KEYS = ["be no", "bill of entry"]

_RATE_RE = re.compile(r"(-?\d+(?:,\d{3})*(?:\.\d+)?)")

MAX_ROWS = 50_000
MAX_SHEETS = 5
XLSX_MAGIC_PREFIX = b"PK\x03\x04"
EXCEL_CELL_MAX = 32_000


def _find_col(df: pd.DataFrame, keys: list[str]) -> str | None:
    cols_lower = {str(c).strip().lower(): c for c in df.columns}
    for k in keys:
        for low, real in cols_lower.items():
            if k in low:
                return real
    return None


def _parse_num(v) -> float | None:
    if v is None or pd.isna(v):
        return None
    match = _RATE_RE.search(str(v))
    if not match:
        return None
    value = float(match.group(1).replace(",", ""))
    return value if math.isfinite(value) and value >= 0 else None


def load_be_excel(path: pathlib.Path) -> tuple[pd.DataFrame, dict]:
    """Load first sheet with read-only streaming, enforcing row/sheet limits."""
    suffix = path.suffix.lower()
    if suffix != ".xlsx":
        raise ValueError("Only .xlsx workbooks are supported (not .xls / .xlsm).")
    with path.open("rb") as stream:
        head = stream.read(4)
    if head != XLSX_MAGIC_PREFIX:
        raise ValueError("Not a valid .xlsx (ZIP) file.")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise ValueError(f"Too many worksheets (max {MAX_SHEETS}).")
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise ValueError("Empty workbook.")
        headers = [str(c).strip() if c is not None else "" for c in header]
        data: list[tuple] = []
        for i, row in enumerate(rows_iter):
            if i >= MAX_ROWS:
                raise ValueError(f"Too many rows (max {MAX_ROWS}).")
            data.append(row)
    finally:
        wb.close()

    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("Every column needs a unique, nonempty header")
    df = pd.DataFrame(data, columns=headers)
    cols = {
        "description": _find_col(df, DESC_KEYS),
        "cth": _find_col(df, CTH_KEYS),
        "bcd": _find_col(df, BCD_KEYS),
        "ntfn": _find_col(df, NTFN_KEYS),
        "assess_val": _find_col(df, ASSESS_KEYS),
        "importer": _find_col(df, IMPORTER_KEYS),
        "be_no": _find_col(df, BE_NO_KEYS),
    }
    if not cols["description"]:
        raise ValueError("Could not find item description column. Expected one of: " + ", ".join(DESC_KEYS))
    if not cols["cth"]:
        raise ValueError("Could not find CTH column. Expected one of: " + ", ".join(CTH_KEYS))
    return df, cols


def _to_cth_str(v) -> str:
    """Remove spreadsheet formatting without turning six digits into eight."""
    if v is None or pd.isna(v):
        return ""
    value = str(v).strip()
    value = re.sub(r"\.0+$", "", value)
    compact = re.sub(r"[ .-]", "", value)
    return compact if compact.isdigit() else value


def _trim_cell(val, max_len: int = EXCEL_CELL_MAX) -> str:
    if val is None or val == "":
        return ""
    t = str(val)
    return t if len(t) <= max_len else t[:max_len]


def normalise_input_rows(df: pd.DataFrame, cols: dict) -> list[dict]:
    rows: list[dict] = []
    df = df.fillna("")
    for i, r in df.iterrows():
        rows.append({
            "row_idx": int(i),
            "item_description": str(r.get(cols["description"], "") or "").strip(),
            "declared_cth": _to_cth_str(r.get(cols["cth"], "")),
            "declared_bcd": _parse_num(r.get(cols["bcd"]) if cols["bcd"] else None),
            "ntfn_ref": str(r.get(cols["ntfn"], "") or "").strip() if cols["ntfn"] else "",
            "assess_val": _parse_num(r.get(cols["assess_val"]) if cols["assess_val"] else None),
            "importer": str(r.get(cols["importer"], "") or "").strip() if cols["importer"] else "",
            "be_no": str(r.get(cols["be_no"], "") or "").strip() if cols["be_no"] else "",
        })
    return rows


# --- Annotated output --------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FRAUD_UNDER_FILL = PatternFill("solid", fgColor="FFD7D7")
BORDER_RED = Border(left=Side(style="thick", color="C00000"))


def write_annotated_excel(
    input_df: pd.DataFrame,
    results: list[dict],
    output_path: pathlib.Path,
) -> None:
    """Write enriched Excel with original cols + classification + fraud columns,
    formatted for auditors (frozen header, autofilter, fraud row highlighting)."""
    by_idx = {r["row_idx"]: r for r in results}
    extras: dict[str, list] = {
        "Suggested CTH": [], "Accepted CTH": [], "Reference BCD %": [], "Method": [], "Model score": [],
        "Fraud Flag": [], "Fraud Type": [], "Revenue at Risk": [],
        "Reasoning": [], "Disagreement": [], "Applicable Notifications": [],
        "Review flags": [], "BCD basis": [],
        "Comments": [],
    }
    for i in range(len(input_df)):
        r = by_idx.get(i)
        if r is None:
            for k in extras:
                extras[k].append("")
            continue
        extras["Suggested CTH"].append(r["final"].get("suggested_cth") or r["final"]["cth_final"])
        extras["Accepted CTH"].append(r["final"]["cth_final"])
        extras["Reference BCD %"].append(r["final"]["bcd_final"] if r["final"]["bcd_final"] is not None else "")
        extras["Method"].append(r["final"]["tier_used"])
        extras["Model score"].append(round(r["final"]["confidence"], 3))
        extras["Fraud Flag"].append("YES" if r["fraud"]["is_fraud"] else "")
        extras["Fraud Type"].append(r["fraud"]["flag_type"])
        extras["Revenue at Risk"].append(
            r["fraud"]["revenue_at_risk"] if r["fraud"]["revenue_at_risk"] is not None else ""
        )
        extras["Reasoning"].append(_trim_cell(r["final"]["reasoning"]))
        extras["Disagreement"].append(_trim_cell(r["final"]["disagreement"]))
        extras["Applicable Notifications"].append(", ".join(r.get("applicable_notifications", [])))
        extras["Review flags"].append(", ".join(r.get("review_flags", [])))
        rr = r.get("rate_resolution") or {}
        extras["BCD basis"].append(str(rr.get("bcd_basis", "") or ""))
        extras["Comments"].append(_trim_cell(_format_audit_comment(r)))

    base = input_df.reset_index(drop=True).astype(object).fillna("")
    extras_df = pd.DataFrame(extras, dtype=object)
    df = pd.concat([base, extras_df], axis=1)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Audit")
        ws = w.sheets["Audit"]

        for col_idx in range(1, len(df.columns) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.freeze_panes = "A2"
        if df.shape[0]:
            ws.auto_filter.ref = ws.dimensions

        fraud_col_idx = list(df.columns).index("Fraud Flag") + 1
        flag_col_idx = list(df.columns).index("Fraud Type") + 1
        for r_idx in range(2, len(df) + 2):
            is_fraud = (ws.cell(row=r_idx, column=fraud_col_idx).value or "") == "YES"
            ftype = ws.cell(row=r_idx, column=flag_col_idx).value or ""
            if is_fraud and ("Under-declared" in ftype or ftype == "Both"):
                for c_idx in range(1, len(df.columns) + 1):
                    ws.cell(row=r_idx, column=c_idx).fill = FRAUD_UNDER_FILL
                ws.cell(row=r_idx, column=1).border = BORDER_RED

        for c_idx, name in enumerate(df.columns, start=1):
            max_len = max([len(str(name))] + [len(str(v)[:50]) for v in df[name].head(200)])
            ws.column_dimensions[get_column_letter(c_idx)].width = min(max(12, max_len), 60)


def _format_audit_comment(row: dict) -> str:
    audit = row.get("audit_trace") or {}
    summary = audit.get("summary") or ""
    bits: list[str] = []
    if summary:
        bits.append(summary)
    rr = row.get("rate_resolution") or {}
    if rr.get("detail"):
        bits.append(f"- bcd_resolution: {rr['detail']}")
    rev = row.get("review_flags") or []
    if rev:
        bits.append(f"- review_flags: {', '.join(rev)}")
    for stage in audit.get("stages", []) or []:
        explanation = stage.get("explanation")
        if explanation:
            bits.append(f"- {stage.get('stage', 'step')}: {explanation}")
    return "\n".join(bits)
