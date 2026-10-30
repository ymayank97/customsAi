import pandas as pd
import pytest
from openpyxl import load_workbook

from backend.app.services.bcd_assessment import assess, attach
from backend.app.services.excel_processor import write_annotated_excel


def result(provided=10, reference=10):
    return {"row_idx": 0, "declared_cth": "73181500", "declared_bcd": provided,
            "item_description": "steel bolts", "review_flags": [],
            "final": {"cth_final": "73181500", "bcd_final": reference, "manual_queue": False,
                      "tier_used": "hierarchical_model", "confidence": 0.9,
                      "reasoning": "test", "disagreement": ""},
            "rate_resolution": {"bcd_basis": "statutory"},
            "fraud": {"is_fraud": False, "flag_type": "", "revenue_at_risk": None}}


@pytest.mark.parametrize("provided,reference,score,status,priority,underpayment", [
    (10, 10, 100, "match", "Low", False),
    (0, 0, 100, "match", "Low", False),
    (5, 10, 50, "under_declared", "High", True),
    (0, 10, 0, "under_declared", "High", True),
    (10, 5, 50, "over_declared", "Low", False),
    (10, 0, 0, "over_declared", "Low", False),
    (1.01, 1, 100, "match", "Low", False),
])
def test_rate_agreement_and_direction(provided, reference, score, status, priority, underpayment):
    assessment = assess(result(provided, reference))
    assert assessment["score"] == score
    assert assessment["status"] == status
    assert assessment["fraud_review_priority"] == priority
    assert assessment["potential_underpayment"] is underpayment
    assert assessment["requires_review"] is (status != "match")


@pytest.mark.parametrize("provided,reference", [(None, 10), (10, None), (float('nan'), 10), (10, float('inf')), (-1, 10)])
def test_missing_invalid_or_nonfinite_rates_are_not_scored(provided, reference):
    assessment = assess(result(provided, reference))
    assert assessment["score"] is None
    assert assessment["status"] == "unverified"
    assert assessment["requires_review"]
    assert assessment["potential_underpayment"] is None
    assert assessment["fraud_review_priority"] == "Review required"


def test_matching_numbers_do_not_hide_uncertain_classification_or_conditions():
    row = result()
    row["final"]["manual_queue"] = True
    assert assess(row)["score"] is None
    row = result()
    row["rate_resolution"]["bcd_basis"] = "manual_review"
    assert assess(row)["score"] is None
    row = result()
    row["review_flags"] = ["rate_requires_review"]
    assert assess(row)["score"] is None


def test_cth_mismatch_remains_red_even_when_bcd_matches():
    row = result()
    row["declared_cth"] = "52081100"
    assessment = assess(row)
    assert assessment["score"] == 100
    assert assessment["requires_review"]
    assert assessment["fraud_review_priority"] == "Medium"
    assert assessment["potential_underpayment"] is False


def test_export_includes_priority_and_highlights_both_incorrect_and_uncertain_rows(tmp_path, monkeypatch):
    monkeypatch.setattr("backend.app.services.classifier.get_master_row", lambda code: None)
    rows = [result(10, 10), result(5, 10), result(10, 5), result(0, 0)]
    rows[-1]["rate_resolution"]["bcd_basis"] = "manual_review"
    for i, row in enumerate(rows):
        row["row_idx"] = i
    target = tmp_path / "audit.xlsx"
    write_annotated_excel(pd.DataFrame({"Description": ["bolt"] * 4}), rows, target)
    workbook = load_workbook(target)
    sheet = workbook["Audit"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["BCD match score /100"]).value == 100
    assert sheet.cell(3, headers["BCD match score /100"]).value == 50
    assert sheet.cell(3, headers["Fraud review priority"]).value == "High"
    assert sheet.cell(4, headers["Fraud review priority"]).value == "Low"
    assert sheet.cell(5, headers["BCD match score /100"]).value is None
    assert sheet.cell(5, headers["Potential BCD underpayment"]).value == "UNKNOWN"
    assert sheet.cell(2, 1).fill.fill_type is None
    for row in (3, 4, 5):
        assert all(cell.fill.fgColor.rgb == "00FFD7D7" for cell in sheet[row])
    assert sheet.cell(1, headers["BCD match score /100"]).comment
    workbook.close()


def test_estimate_shows_candidate_and_rate_agreement_without_accepting_it(monkeypatch):
    monkeypatch.setattr("backend.app.services.classifier.get_master_row", lambda code: {"bcd_statutory": 10})
    row = result(5)
    row["final"].update(cth_final="", bcd_final=None, manual_queue=True,
                        suggested_cth="85", candidates=[{"cth": "85366990"}])
    enriched = attach(row)
    assert enriched["prediction_estimate"]["cth"] == "85366990"
    assert enriched["prediction_estimate"]["bcd"] == 10
    assessment = enriched["bcd_assessment"]
    assert assessment["score"] == 50
    assert assessment["status"] == "estimated"
    assert not assessment["verified"]
    assert assessment["requires_review"]
    assert assessment["potential_underpayment"] is None
    assert enriched["final"]["cth_final"] == ""
    assert row.get("prediction_estimate") is None
    row["review_flags"] = ["legal_rule_conflict"]
    assert attach(row)["prediction_estimate"]["cth"] == ""
    assert attach(row)["bcd_assessment"]["score"] is None


def test_no_candidate_never_copies_declared_code_or_rate(monkeypatch):
    monkeypatch.setattr("backend.app.services.classifier.get_master_row", lambda code: None)
    row = result()
    row["final"].update(cth_final="", bcd_final=None, manual_queue=True, candidates=[])
    enriched = attach(row)
    assert enriched["prediction_estimate"]["cth"] == ""
    assert enriched["bcd_assessment"]["score"] is None
